# -*- coding: utf-8 -*-
"""Build the R11 adjudication kit (review item 2).
  Sheet ① independent re-annotation of all 120 repositories under the codebook (original judgments hidden)
  Sheet ② consensus adjudication of the 40 non-unanimous repositories (original judgments shown)
Input: ../01_reanalysis/_human120.pkl (written by r11_reanalysis.py), the short descriptions shown to annotators (originally Korean, English translation in this repository), raw metadata.
Output: R11_adjudication_sheets.xlsx
"""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
HERE = os.path.dirname(os.path.abspath(__file__))
if os.path.isdir(os.path.join(os.path.dirname(HERE), 'results')):     # replication repository: audit/<this file>
    REPO = os.path.dirname(HERE); AUD = HERE; DATA = os.path.join(REPO, 'data'); PKL = os.path.join(REPO, 'results', 'R11', '_human120.pkl')
else:                                                                 # authors' working folder: 92_…/02_…/<this file>
    ROOT = os.path.dirname(os.path.dirname(HERE)); AUD = os.path.join(ROOT, '40_audit_human_annotation_kit'); DATA = os.path.join(ROOT, '10_data_collection')
    PKL = os.path.join(HERE, '..', '01_reanalysis', '_human120.pkl')
H = pd.read_pickle(PKL)
ko = pd.read_excel(os.path.join(AUD, 'survey_simplest_completed.xlsx'), sheet_name='① Projects (120)').iloc[:, :2]; ko.columns = ['item_no', 'desc_ko']
raw = pd.read_excel(os.path.join(DATA, 'github_ai_agent_raw_v2.xlsx'))[['full_name', 'description', 'topics']]
H = H.merge(ko, on='item_no').merge(raw, on='full_name', how='left')
CLS = {'1': '1 Agent software', '2': '2 Learning/reference resource', '3': '3 Unrelated/undecidable'}
YEL = PatternFill('solid', fgColor='FFF2CC'); HDR = PatternFill('solid', fgColor='D9E1F2'); WRAP = Alignment(wrap_text=True, vertical='top')
wb = Workbook(); ws0 = wb.active; ws0.title = 'Read me'
for i, line in enumerate([
    'R11 adjudication kit: read the codebook (R11_codebook.md, adjudication criteria) first.',
    'Sheet ① (independent re-annotation): all 120 items. The original judgments are hidden. Do not discuss with anyone. Enter 1/2/3 and the rule number (R1-R10) in the yellow cells.',
    'Sheet ② (consensus adjudication): the 40 items on which the original three annotators disagreed. Look at the original judgments and decide the final class based on the codebook rules.',
    'Sheet ①: each rater fills in a separate copy of the file (file name ending in _rater1, _rater2). Sheet ②: one file, by a single adjudicator or by consensus of two.',
    'Basis for judgment: repository name, description and tags first; if insufficient, look only at the first screen of the README via the link.'], 1):
    ws0.cell(i, 1, line).alignment = WRAP
ws0.column_dimensions['A'].width = 120
def sheet(title, rows, show_orig):
    ws = wb.create_sheet(title)
    cols = ['item_no', 'repository', 'Link', 'Short description', 'description (original)', 'topics'] + (['A original judgment', 'B original judgment', 'C original judgment', 'Votes for agent'] if show_orig else []) + ['Judgment (1/2/3)', 'Rule no.', 'Notes']
    for j, c in enumerate(cols, 1):
        x = ws.cell(1, j, c); x.fill = HDR; x.font = Font(bold=True); x.alignment = WRAP
    dv = DataValidation(type='list', formula1='"1,2,3"', allow_blank=True); ws.add_data_validation(dv)
    for i, r in enumerate(rows.itertuples(), 2):
        vals = [r.item_no, r.full_name, f'https://github.com/{r.full_name}', r.desc_ko, r.description, r.topics]
        if show_orig: vals += [CLS[r.A_raw], CLS[r.B_raw], CLS[r.C_raw], int(r.n_pos)]
        vals += [None, None, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.alignment = WRAP
        jcls = len(vals) - 2
        for j in (jcls, jcls + 1, jcls + 2): ws.cell(i, j).fill = YEL
        dv.add(ws.cell(i, jcls))
    for j, w in enumerate([6, 30, 32, 50, 50, 28] + ([16, 16, 16, 7] if show_orig else []) + [11, 10, 30], 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    ws.freeze_panes = 'B2'
# original 3-class answers (the pickle keeps only binary columns), read with the scoring rule of score_three_annotators.py
MARK = lambda v: str(v).strip() not in ('', 'nan', 'None') and str(v).strip().upper().rstrip('.0') in ('O', '', 'V', 'X', '✓', '\u3147', '○', 'Ο')  # accepted check marks; '\u3147' is the Korean letter ieung, which looks like O
FILES = {'A': 'survey_simplest_completed.xlsx', 'B': 'survey_friend_version_rater2_completed.xlsx', 'C': 'survey_friend_version_rater3_completed.xlsx'}
for k, f in FILES.items():
    d = pd.read_excel(os.path.join(AUD, f), sheet_name='① Projects (120)'); d.columns = ['item_no', 'desc', 'link', 'a1', 'a2', 'a3']; d['item_no'] = range(1, len(d) + 1)
    d[k + '_raw'] = d.apply(lambda r: str([MARK(r.a1), MARK(r.a2), MARK(r.a3)].index(True) + 1), axis=1)
    H = H.merge(d[['item_no', k + '_raw']], on='item_no')
assert all(((H[k + '_raw'] == '1') == H[k + '_pos']).all() for k in 'ABC'),'binary columns disagree with raw answers'
sheet('① Independent reannotation 120', H.sample(frac=1, random_state=2026)[['item_no', 'full_name', 'desc_ko', 'description', 'topics']], False)
sheet('② Consensus adjudication 40', H[H['disagree']].sort_values('item_no'), True)
out = os.path.join(HERE, 'R11_adjudication_sheets.xlsx'); wb.save(out)
print('saved', out, '| sheet1 =', len(H), '| sheet2 =', int(H['disagree'].sum()))
