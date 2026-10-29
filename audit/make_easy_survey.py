# -*- coding: utf-8 -*-
"""Generate a very easy single-annotator questionnaire (2026-09-10)
  survey_easy_version.xlsx  - the file the annotator opens. No model labels (blind).
     Sheet 1 'Repositories (120)' : no., name, description, tags, link + [Judgment] dropdown with 3 options
     Sheet 2 'Cluster labels (17)'  : terms, representative repos, candidate label + [Does it fit?] dropdown with 3 options
  survey_answer_key_do_not_open.xlsx - LLM labels and cluster IDs. Used for scoring only.
Scoring: python score_easy_survey.py
"""
import os, random, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
random.seed(2026)
HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(os.path.dirname(HERE), 'results') if os.path.isdir(os.path.join(os.path.dirname(HERE), 'results')) else os.path.join(os.path.dirname(HERE), '30_analysis_results')
PER_CLUSTER = 24

# -- 1. 120 repositories: 24 per segment, stratified from the 404-repository audit sample --
aud = pd.read_excel(os.path.join(HERE, 'human_eval_stratified_sample_400.xlsx'), sheet_name='sample')
picks = []
for c in sorted(aud['cluster'].unique()):
    g = aud[aud['cluster'] == c]
    picks.append(g.sample(min(PER_CLUSTER, len(g)), random_state=100 + int(c)))
s = pd.concat(picks).sample(frac=1, random_state=7).reset_index(drop=True)
s.insert(0, 'item_no', range(1, len(s) + 1))

# -- 2. 17 cluster labels --
mac = pd.read_excel(os.path.join(RES, 'topic_macro_k5.xlsx')); mac['cid'] = 'T' + mac['cluster'].astype(str)
mes = pd.read_excel(os.path.join(RES, 'topic_meso_k12.xlsx')); mes['cid'] = 'M' + mes['cluster'].astype(str)
lab = pd.concat([mac, mes], ignore_index=True).sample(frac=1, random_state=3).reset_index(drop=True)
lab.insert(0, 'item_no', range(1, len(lab) + 1))

# -- 3. Questionnaire xlsx --
wb = Workbook(); ws = wb.active; ws.title = 'Repositories (120)'
head_fill = PatternFill('solid', fgColor='DDEBF7'); ans_fill = PatternFill('solid', fgColor='FFF2CC')
ws.append(['item_no', 'Repository name', 'Description', 'Tags', 'Link (click)', '▶ Judgment (choose from dropdown)', 'Notes (optional)'])
for _, r in s.iterrows():
    ws.append([int(r['item_no']), r['full_name'], str(r['description']), str(r['topics']).replace(';', ', '), r['url'], '', ''])
    ws.cell(ws.max_row, 5).hyperlink = r['url']; ws.cell(ws.max_row, 5).font = Font(color='0563C1', underline='single')
dv = DataValidation(type='list', formula1='"1 Yes - AI agent,2 Tutorial/resource list/paper code,3 No or not sure"', allow_blank=True, showDropDown=False)
dv.error = 'Please choose from the dropdown'; ws.add_data_validation(dv); dv.add(f'F2:F{ws.max_row}')
for col, w in zip('ABCDEFG', [6, 34, 70, 40, 14, 34, 24]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
for c in ws[1]: c.fill = head_fill; c.font = Font(bold=True)
for i in range(2, ws.max_row + 1): ws.cell(i, 6).fill = ans_fill
ws.freeze_panes = 'A2'

ws2 = wb.create_sheet('Cluster labels (17)')
ws2.append(['item_no', 'Representative terms of this cluster', 'Representative repositories', 'Name (label) given by AI', '▶ Does the name fit? (dropdown)', '▶ Do the terms form one topic? (dropdown)', 'Notes (optional)'])
for _, r in lab.iterrows():
    ws2.append([int(r['item_no']), r['top_terms'], str(r['rep_repos']).replace(';', '\n'), r['label_refined'], '', '', ''])
dv2 = DataValidation(type='list', formula1='"3 Fits well,2 Roughly fits,1 Does not fit"', allow_blank=True); ws2.add_data_validation(dv2); dv2.add(f'E2:E{ws2.max_row}'); dv2.add(f'F2:F{ws2.max_row}')
for col, w in zip('ABCDEFG', [6, 50, 44, 40, 26, 30, 24]): ws2.column_dimensions[col].width = w
for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
for c in ws2[1]: c.fill = head_fill; c.font = Font(bold=True)
for i in range(2, ws2.max_row + 1): ws2.cell(i, 5).fill = ans_fill; ws2.cell(i, 6).fill = ans_fill
ws2.freeze_panes = 'A2'

ws0 = wb.create_sheet('Read me', 0)
for line in ['This questionnaire is completed by one person in under an hour.', '',
             '[Sheet 1] Repositories (120): look at the name, description and tags and choose exactly one option in the yellow cell.',
             '   1 Yes - AI agent   = tools/frameworks for building agents, apps that operate as agents, MCP servers',
             '   2 Tutorial/resource list/paper code = course examples, awesome lists, paper reproduction code',
             '   3 No or not sure = unrelated to AI agents, or cannot be judged',
             '   * If the description alone is ambiguous, click the link and look for 10 seconds. Do not deliberate for more than 20 seconds.',
             '', '[Sheet 2] Cluster labels (17): look at the term set and representative repos, and rate whether the AI-given name fits (3/2/1) and whether the terms form one topic (3/2/1).',
             '', 'When finished, save the file and let us know. Scoring is automatic.']:
    ws0.append([line])
ws0.column_dimensions['A'].width = 110
out = os.path.join(HERE, 'survey_easy_version.xlsx'); wb.save(out)

# -- 4. Answer key (separate file) --
with pd.ExcelWriter(os.path.join(HERE, 'survey_answer_key_do_not_open.xlsx')) as w:
    s[['item_no', 'full_name', 'cluster', 'year', 'llm_qwen', 'llm_llama']].to_excel(w, sheet_name='repo_key', index=False)
    lab[['item_no', 'cid', 'n_docs', 'label_refined']].to_excel(w, sheet_name='label_key', index=False)
print('Questionnaire with 120+17 items created:', out)
