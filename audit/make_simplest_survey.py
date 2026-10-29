# -*- coding: utf-8 -*-
"""Simplest questionnaire: short descriptions only, large font, just type one O. -> survey_simplest.xlsx (scoring: score_simplest.py)"""
import os, json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
HERE = os.path.dirname(os.path.abspath(__file__))
tr = json.load(open(os.path.join(HERE, 'survey_annotator_display_text.json'), encoding='utf-8'))
base = pd.read_excel(os.path.join(HERE, 'survey_easy_version_english_only.xlsx'), sheet_name='Repositories (120)')
labs = pd.read_excel(os.path.join(HERE, 'survey_easy_version_english_only.xlsx'), sheet_name='Cluster labels (17)')
thin = Side(style='thin', color='BBBBBB'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
H = PatternFill('solid', fgColor='305496'); Y = PatternFill('solid', fgColor='FFF2CC'); G = PatternFill('solid', fgColor='E2EFDA')
big = Font(size=13); bold = Font(size=13, bold=True, color='FFFFFF')

wb = Workbook(); ws = wb.active; ws.title = '① Projects (120)'
ws.append(['item_no', 'What is this project? (short description)', 'Link', 'AI agent\nyes', 'Resource list/\ntutorial', 'No/\nnot sure'])
for _, r in base.iterrows():
    ws.append([int(r['item_no']), tr['repo'].get(str(r['item_no']), ''), 'Open', '', '', ''])
    c = ws.cell(ws.max_row, 3); c.hyperlink = r['Link (click)']; c.font = Font(size=12, color='0563C1', underline='single')
for col, w in zip('ABCDEF', [6, 72, 7, 13, 13, 13]): ws.column_dimensions[col].width = w
ws.row_dimensions[1].height = 42
for r in range(1, ws.max_row + 1):
    ws.row_dimensions[r].height = ws.row_dimensions[r].height or 48
    for c in range(1, 7):
        cell = ws.cell(r, c); cell.border = box; cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center' if c != 2 else 'left')
        if r == 1: cell.fill = H; cell.font = bold
        else:
            if c != 3: cell.font = big
            if c >= 4: cell.fill = Y
            if c == 2: cell.fill = G
ws.freeze_panes = 'A2'

ws2 = wb.create_sheet('② Group names (17)')
ws2.append(['item_no', 'Main terms of this group', 'Group name given by AI', 'Name\nfits well', 'Roughly\nfits', 'Does not fit'])
for _, r in labs.iterrows():
    terms = [t.strip() for t in str(r['Representative terms of this cluster']).split(',')][:10]
    ws2.append([int(r['item_no']), ', '.join(tr['terms'].get(t, t) for t in terms), tr['label'].get(str(r['item_no']), ''), '', '', ''])
for col, w in zip('ABCDEF', [6, 58, 34, 12, 12, 12]): ws2.column_dimensions[col].width = w
ws2.row_dimensions[1].height = 42
for r in range(1, ws2.max_row + 1):
    if r > 1: ws2.row_dimensions[r].height = 52
    for c in range(1, 7):
        cell = ws2.cell(r, c); cell.border = box; cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center' if c not in (2, 3) else 'left')
        if r == 1: cell.fill = H; cell.font = bold
        else:
            cell.font = big
            if c >= 4: cell.fill = Y
            if c in (2, 3): cell.fill = G
ws2.freeze_panes = 'A2'

ws0 = wb.create_sheet('Read before starting', 0)
lines = ['How to do it (about 1 hour in total)', '',
         '① Projects (120) sheet: read the short description and type O in exactly one of the three yellow cells.',
         '     AI agent yes = tools/frameworks for building agents, apps that operate as agents, MCP servers',
         '     Resource list/tutorial = link collections (awesome), lectures/examples, paper code',
         '     No/not sure = unrelated to AI agents, or cannot be judged',
         '     * If the description is ambiguous, click "Open" and look for 10 seconds. Do not deliberate for more than 20 seconds.', '',
         '② Group names (17) sheet: look at the terms and type O in one of the three cells for whether the AI-given name fits.', '',
         'When finished, save (Ctrl+S) and let us know. Scoring is automatic.']
for l in lines: ws0.append([l])
ws0.column_dimensions['A'].width = 100
for r in range(1, ws0.max_row + 1): ws0.cell(r, 1).font = Font(size=13, bold=(r == 1))
out = os.path.join(HERE, 'survey_simplest.xlsx'); wb.save(out); print('saved', out)
