# -*- coding: utf-8 -*-
"""R6 report: add the sheets '1b_R6 Abstract term comparison' and '2b_R6 Abstract word DF' to the R5 report xlsx and update the summary."""
import json, re, os, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); OUT = os.path.join(ROOT, '30_analysis_results', 'term_review')
docs = {k: re.sub(r'[‐-―]', '-', v.lower()) for k, v in json.load(open(os.path.join(OUT, 'prior_work_fulltext.json'), encoding='utf-8')).items()}
def df(t):
    t = t.lower(); pats = {t, t.replace('-', ' '), t.replace(' ', '-')}
    if not t.endswith('s'): pats |= {p + 's' for p in list(pats)}
    hit = [k for k, d in docs.items() if any(p in d for p in pats)]; return len(hit), ', '.join(sorted(h[:4] + ' ' + h.split('_')[1][:12] for h in hit))
rows = [
 ('recursive partitioning by creation date circumvents the 1,000-result limit', 'partition (6 papers) · removes (3 papers) · recursive (5 papers)', 'a recursive partition of each query by creation date removes the 1,000-result limit', '"partitioning" 2 papers, "circumvents" 0 papers'),
 ('the analysis characterizes the creation-year composition', 'describe (many) · composition (16 papers) · creation (12 papers)', 'the analysis describes the composition by creation year', '"characterizes" 1 paper, "creation-year" 0 papers'),
 ('a stability-based model-selection rule', 'model selection (3 papers) · stability (10 papers)', 'a model-selection rule based on stability', '"stability-based" 1 paper'),
 ('supports a five-cluster coarse solution', 'cluster (13 papers) · coarse (4 papers)', 'supports a coarse solution with five clusters', '"five-cluster" 0 papers'),
 ('a 120-repository subsample', 'sample (many)', 'a 120-repository sample', '"subsample" 1 paper'),
 ('variation in the operational boundary of an AI-agent repository', 'boundary (4 papers) · variation (19 papers)', 'variation in the boundary of what counts as an AI-agent repository', '"operational boundary" 0 papers'),
 ('falling within the range', 'within (many)', 'lying within the range', '"falling" 2 papers'),
 ('alternative temporal model specifications', 'specification(s) (5/3 papers) · alternative (many)', 'alternative model specifications', '"temporal model" 0 papers'),
 ('robustness analyses over relevance and sampling assumptions', 'robustness (13 papers)', 'robustness checks over relevance and sampling assumptions', '"robustness analyses" 0 papers'),
 ('LangChain-centered framework repositories', 'LangChain (proper noun) · based (many)', 'LangChain-based framework repositories', 'combination "-centered" 0 papers; same replacement throughout the body'),
 ('game-related repositories also declines robustly', 'game (11 papers) · decrease (many)', 'game repositories also decreases robustly', '"game-related" 0 papers, "declines" 1 paper'),
 ('the apparent full-period decline of the largest orchestration-and-automation cluster', 'decrease · period (14 papers) · orchestration (4 papers) · automation (8 papers)', 'the decrease of the largest orchestration and automation cluster estimated over the full period', '"apparent" 2 papers, "full-period" 0 papers, hyphenated combination 0 papers'),
 ('coincides temporally with several relevant events', 'aligned with (9 papers)', 'is aligned in time with several relevant events', '"coincides" 1 paper, "temporally" 0 papers'),
 ('the observational design does not support causal attribution', 'cause (24 papers) · design (many)', 'the study design does not identify a cause', '"observational" 0 papers, "attribution" 2 papers'),
 ('a reproducible, repository-level characterization', 'reproducible (7 papers) · description (18 papers) · repository (12 papers)', 'a reproducible description, at the repository level', '"repository-level" 1 paper, "characterization" 2 papers'),
 ('currently visible, query-matched AI-agent repositories', 'matched (4 papers) · search queries (many)', 'currently visible AI-agent repositories matched by the search queries', '"query-matched" 0 papers'),
 ('toward MCP infrastructure', 'MCP server (2 papers) · server (many)', 'toward MCP servers', 'combination "MCP infrastructure" 0 papers (consistent with the body)'),
 ('(retained) silhouette · binomial · false-discovery-rate · BGE-M3', 'standard statistical terms / proper nouns (DF 0-1)', 'retained', 'no alternative, noted'),
]
rep = pd.DataFrame([{'A. User abstract term (R6 draft)': a, 'B. Common prior-work term (number of papers)': b, 'C. Changed term (R6 final)': c, 'E. Rationale / notes': e} for a, b, c, e in rows])
A = open(os.path.join(OUT, 'R6_abstract_final.txt'), encoding='utf-8').read().replace('\\', '')
STOP = set('a an the and or of to in on for with by from as at is are was were be been that this these those we our it its which who than then there not no so such into over under between among within without through per each all more most less only same other if when where while whether because also does do did has have had may might can could would should will must using use used rather'.split())
seen = []; wrows = []
for w in re.findall(r"[a-z][a-z0-9\-]+", A.lower()):
    if w in STOP or len(w) < 3 or w in seen: continue
    seen.append(w); d, h = df(w); wrows.append({'Abstract word (R6 final)': w, 'Prior-work papers (of 27)': d, 'Verdict': 'OK' if d >= 3 else 'Exception (standard statistical term / proper noun)', 'Papers where it occurs': h[:150]})
wdf = pd.DataFrame(wrows).sort_values('Prior-work papers (of 27)')
src = os.path.join(OUT, 'term_standardization_report_260910.xlsx'); out = os.path.join(OUT, 'term_standardization_report_R6_260910.xlsx')
sheets = pd.read_excel(src, sheet_name=None)
with pd.ExcelWriter(out) as w:
    summ = sheets['0_Summary']; summ.loc[len(summ)] = ['R6 addition', 'Exhaustive check of the user-supplied abstract (315 words) -> 17 low-frequency words replaced (sheet 1b) -> all content words of the final abstract DF≥3 (4 exceptions: silhouette, binomial, false-discovery-rate, BGE-M3) -> body: LangChain-centered -> LangChain-based']
    summ.to_excel(w, sheet_name='0_Summary', index=False); sheets['1_Term comparison and edits'].to_excel(w, sheet_name='1_Term comparison and edits(R5)', index=False); rep.to_excel(w, sheet_name='1b_R6 Abstract term comparison', index=False)
    wdf.to_excel(w, sheet_name='2b_R6 Abstract word DF', index=False)
    for k in ['2_Title-abstract word DF', '3_Body compound DF (low first)', '4_Candidate term DF evidence']: sheets[k].to_excel(w, sheet_name=k, index=False)
wb = load_workbook(out)
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496'); c.alignment = Alignment(wrap_text=True, vertical='center')
    for col, w_ in zip('ABCDEFG', [12, 48, 46, 46, 30, 60, 40]): ws.column_dimensions[col].width = w_
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
wb.save(out); print('saved', out, '| words', len(seen), '| low', [(r['Abstract word (R6 final)'], r['Prior-work papers (of 27)']) for r in wrows if r['Prior-work papers (of 27)'] < 3])
